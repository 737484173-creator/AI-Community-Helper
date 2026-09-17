from __future__ import annotations

import argparse
from dataclasses import dataclass
from pathlib import Path
from typing import Any

from openpyxl import load_workbook

from app.database.migrations.init_db import init_database
from app.database.repositories.faq_repository import FAQRepository


REQUIRED_FIELDS = ("category", "question", "answer")
OPTIONAL_FIELDS = ("subcategory", "keywords", "enabled")
SUPPORTED_FIELDS = REQUIRED_FIELDS + OPTIONAL_FIELDS


@dataclass(frozen=True)
class FAQImportResult:
    imported_count: int
    skipped_count: int
    categories: list[str]


def _normalize_field_name(value: Any) -> str:
    return str(value or "").strip().lower().replace(" ", "_")


def _strip_prefixed_value(value: Any, expected_field: str) -> str:
    text = str(value or "").strip()
    prefix = f"{expected_field}:"
    if text.lower().startswith(prefix):
        return text[len(prefix) :].strip()
    return text


def _parse_enabled(value: Any) -> bool:
    if value is None or str(value).strip() == "":
        return True
    return str(value).strip().lower() not in {"0", "false", "no", "n", "disabled"}


def _row_from_prefixed_cells(row: tuple[Any, ...]) -> dict[str, Any]:
    values: dict[str, Any] = {}
    for cell in row:
        if cell is None:
            continue
        text = str(cell).strip()
        for field in SUPPORTED_FIELDS:
            prefix = f"{field}:"
            if text.lower().startswith(prefix):
                values[field] = text[len(prefix) :].strip()
                break
    return values


def _row_from_headers(headers: list[str], row: tuple[Any, ...]) -> dict[str, Any]:
    values: dict[str, Any] = {}
    for index, header in enumerate(headers):
        if header in SUPPORTED_FIELDS and index < len(row):
            values[header] = row[index]
    return values


def _is_header_row(row: tuple[Any, ...]) -> bool:
    normalized = {_normalize_field_name(cell) for cell in row if cell is not None}
    return all(field in normalized for field in REQUIRED_FIELDS)


def import_faq_from_excel(
    excel_path: str | Path,
    db_path: str | Path | None = None,
    sheet_name: str | None = None,
) -> FAQImportResult:
    path = Path(excel_path)
    if not path.exists():
        raise FileNotFoundError(f"FAQ Excel file not found: {path}")

    init_database(db_path)
    repository = FAQRepository(db_path)

    workbook = load_workbook(path, read_only=True, data_only=True)
    try:
        worksheet = workbook[sheet_name] if sheet_name else workbook[workbook.sheetnames[0]]
        rows = list(worksheet.iter_rows(values_only=True))
    finally:
        workbook.close()
    if not rows:
        return FAQImportResult(imported_count=0, skipped_count=0, categories=[])

    use_headers = _is_header_row(rows[0])
    headers = [_normalize_field_name(cell) for cell in rows[0]] if use_headers else []
    data_rows = rows[1:] if use_headers else rows

    imported_count = 0
    skipped_count = 0
    categories: set[str] = set()

    for row in data_rows:
        if not any(cell is not None and str(cell).strip() for cell in row):
            skipped_count += 1
            continue

        values = _row_from_headers(headers, row) if use_headers else _row_from_prefixed_cells(row)

        category = _strip_prefixed_value(values.get("category"), "category")
        question = _strip_prefixed_value(values.get("question"), "question")
        answer = _strip_prefixed_value(values.get("answer"), "answer")
        subcategory = _strip_prefixed_value(values.get("subcategory"), "subcategory") or None
        keywords = _strip_prefixed_value(values.get("keywords"), "keywords") or None
        enabled = _parse_enabled(values.get("enabled"))

        if not category or not question or not answer:
            skipped_count += 1
            continue

        repository.upsert(
            category=category,
            subcategory=subcategory,
            question=question,
            answer=answer,
            keywords=keywords,
            enabled=enabled,
        )
        imported_count += 1
        categories.add(category)

    return FAQImportResult(
        imported_count=imported_count,
        skipped_count=skipped_count,
        categories=sorted(categories),
    )


def main() -> None:
    parser = argparse.ArgumentParser(description="Import FAQ knowledge base from Excel.")
    parser.add_argument("excel_path", help="Path to SDK.xlsx or another FAQ workbook.")
    parser.add_argument("--db-path", help="Optional SQLite database path.")
    parser.add_argument("--sheet-name", help="Optional worksheet name.")
    args = parser.parse_args()

    result = import_faq_from_excel(
        excel_path=args.excel_path,
        db_path=args.db_path,
        sheet_name=args.sheet_name,
    )
    print(
        "FAQ import completed: "
        f"imported={result.imported_count}, "
        f"skipped={result.skipped_count}, "
        f"categories={', '.join(result.categories)}"
    )


if __name__ == "__main__":
    main()
