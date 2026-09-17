from __future__ import annotations

from io import BytesIO
from pathlib import Path

from openpyxl import Workbook

from app.dashboard.views import faq_management
from app.database.repositories.faq_repository import FAQRepository
from app.utils.excel_importer import import_faq_from_excel


class UploadedWorkbook:
    def __init__(self, payload: bytes) -> None:
        self._payload = payload

    def getvalue(self) -> bytes:
        return self._payload


class LegacyFAQRepository:
    def list_all(self, limit: int = 200) -> list[object]:
        return [object(), object()][:limit]


def test_importer_supports_prefixed_sdk_rows(tmp_path: Path) -> None:
    workbook = Workbook()
    worksheet = workbook.active
    worksheet.append(
        (
            "Category: Login",
            "Question: Cannot login",
            "Answer: Please restart the game.",
        )
    )
    worksheet.append(
        (
            "Category: Payment",
            "Question: Payment Failed",
            "Answer: Please check your receipt.",
        )
    )

    excel_path = tmp_path / "sdk.xlsx"
    db_path = tmp_path / "community.db"
    workbook.save(excel_path)

    result = import_faq_from_excel(excel_path, db_path)

    assert result.imported_count == 2
    assert result.skipped_count == 0
    assert result.categories == ["Login", "Payment"]

    repository = FAQRepository(db_path)
    login_faqs = repository.list_by_category("Login")
    assert login_faqs[0]["question"] == "Cannot login"
    assert login_faqs[0]["answer"] == "Please restart the game."


def test_importer_supports_header_rows_and_upsert(tmp_path: Path) -> None:
    workbook = Workbook()
    worksheet = workbook.active
    worksheet.append(("Category", "Question", "Answer", "Keywords", "Enabled"))
    worksheet.append(("Login", "Cannot login", "First answer", "login", "yes"))
    worksheet.append(("Login", "Cannot login", "Updated answer", "auth", "yes"))

    excel_path = tmp_path / "faq.xlsx"
    db_path = tmp_path / "community.db"
    workbook.save(excel_path)

    result = import_faq_from_excel(excel_path, db_path)

    assert result.imported_count == 2
    repository = FAQRepository(db_path)
    faq = repository.list_by_category("Login")[0]
    assert faq["answer"] == "Updated answer"
    assert faq["keywords"] == "auth"


def test_dashboard_upload_keeps_a_current_faq_copy(tmp_path: Path, monkeypatch) -> None:
    workbook = Workbook()
    worksheet = workbook.active
    worksheet.append(("Category", "Question", "Answer"))
    worksheet.append(("Login", "Cannot login", "Restart the game."))
    payload = BytesIO()
    workbook.save(payload)

    current_file = tmp_path / "faq" / "current_faq.xlsx"
    db_path = tmp_path / "community.db"
    monkeypatch.setattr(faq_management, "FAQ_SOURCE_FILE", current_file)

    result = faq_management._import_uploaded_file(UploadedWorkbook(payload.getvalue()), db_path)

    assert result.imported_count == 1
    assert current_file.exists()
    assert not current_file.with_name(".current_faq.pending.xlsx").exists()
    assert FAQRepository(db_path).count_all() == 1


def test_faq_count_supports_repository_loaded_before_hot_reload() -> None:
    assert faq_management._faq_record_count(LegacyFAQRepository()) == 2
